# -*- coding: utf-8 -*-
import glob, os
from collections import Counter
import openpyxl

folder = r"C:\Users\admin\Trafic_mojo_2\Eredmények\Teszt_YOLO11S"
files = sorted(glob.glob(os.path.join(folder, "*.xlsx")))

for fpath in files:
    print(f"\n{'='*50}")
    print(f"FILE: {os.path.basename(fpath)}")
    try:
        wb = openpyxl.load_workbook(fpath)
        for shname in wb.sheetnames:
            ws = wb[shname]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(h) if h else '' for h in rows[0]]
            print(f"  Sheet '{shname}' - {len(rows)-1} rows")
            print(f"  Columns: {header}")
            # Find category column
            cat_col = None
            for i, h in enumerate(header):
                if 'ategor' in h.lower() or 'kategor' in h.lower():
                    cat_col = i
                    break
            if cat_col is not None:
                cats = Counter(row[cat_col] for row in rows[1:])
                print(f"  Category breakdown:")
                for k, v in sorted(cats.items(), key=lambda x: -x[1]):
                    print(f"    {k}: {v}")
            else:
                print(f"  (no category column found)")
            # Print all rows for small sheets
            if len(rows) <= 25:
                for i, row in enumerate(rows):
                    print(f"  row{i}: {row}")
    except Exception as e:
        print(f"  Error processing file: {e}")
