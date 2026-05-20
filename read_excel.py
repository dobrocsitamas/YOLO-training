import openpyxl

fpath = r"C:\Users\admin\Trafic_mojo_2\Eredmények\Teszt_YOLO11S\forgalmi_adatok_20260520.xlsx"
wb = openpyxl.load_workbook(fpath)
for shname in wb.sheetnames:
    ws = wb[shname]
    rows = list(ws.iter_rows(values_only=True))
    print(f"\n=== Sheet: {shname} ({len(rows)-1} rows) ===")
    for i, row in enumerate(rows):
        print(f"  {i}: {row}")
