import glob, os
import openpyxl

folder = r"C:\Users\admin\Trafic_mojo_2\Eredmények\Teszt_YOLO11S"
if not os.path.exists(folder):
    print(f"Directory not found: {folder}")
else:
    files = sorted(glob.glob(os.path.join(folder, "*.xlsx")))

    for fpath in files:
        print(f"\n{'='*60}")
        print(f"FILE: {os.path.basename(fpath)}")
        print(f"{'='*60}")
        try:
            wb = openpyxl.load_workbook(fpath)
            for shname in wb.sheetnames:
                ws = wb[shname]
                print(f"\n  Sheet: {shname}")
                rows = list(ws.iter_rows(values_only=True))
                # Print first 5 rows to see structure
                for i, row in enumerate(rows[:5]):
                    print(f"    Row {i+1}: {row}")
                if len(rows) > 5:
                    print(f"    ... ({len(rows)} rows total)")
        except Exception as e:
            print(f"Error reading {fpath}: {e}")
